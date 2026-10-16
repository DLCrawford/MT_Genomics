#!/usr/bin/env python3
"""
04_make_workbook_figure.py
--------------------------
Assemble the six result tables (produced by scripts 01-03) into one formatted
Excel workbook with a README sheet, and render the 3-panel summary figure.

Run 01, 02, 03 first. Outputs:
  selection_stats_bootstrap_CIs.xlsx
  selection_stats_figure.png / .pdf
"""
import csv
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from common import opath

def rd(name):
    with open(opath(name)) as f:
        return [row for row in csv.reader(f, delimiter="\t")]
def val(x):
    try:    return float(x)
    except: return np.nan

# ============================ WORKBOOK ============================
wb = Workbook()
hdrfont  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
hdrfill  = PatternFill("solid", fgColor="2F5496")
titlefont= Font(name="Arial", bold=True, size=13)
notefont = Font(name="Arial", italic=True, size=10, color="444444")
base     = Font(name="Arial", size=11)
thin     = Side(style="thin", color="D9D9D9")
border   = Border(left=thin, right=thin, top=thin, bottom=thin)
center   = Alignment(horizontal="center")

def sheet(ws, title, note, tsv):
    ws["A1"] = title; ws["A1"].font = titlefont
    r = 3
    if note:
        ws.cell(r, 1, note).font = notefont
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 2
    rows = rd(tsv); hdr = rows[0]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(r, j, h); c.font = hdrfont; c.fill = hdrfill; c.border = border; c.alignment = center
    r += 1
    for row in rows[1:]:
        for j, v in enumerate(row, 1):
            try:    v2 = float(v) if "." in v else int(v)
            except: v2 = v
            c = ws.cell(r, j, v2); c.font = base; c.border = border
            if j >= 3: c.alignment = center
            if isinstance(v2, float): c.number_format = "0.000"
        r += 1
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(11, len(str(hdr[j-1])) + 3)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 13)

ws = wb.active; ws.title = "README"
ws["A1"] = "Selection statistics — F. heteroclitus mtDNA (13 protein-coding genes)"
ws["A1"].font = titlefont
readme = [
 "",
 "Three complementary statistics, each with 95% bootstrap CIs (B=5000, site resampling).",
 "",
 "1. pN_pS  — PRIMARY within-species selection statistic (all 141 individuals).",
 "   pN/pS = (obs_N/N_sites)/(obs_S/S_sites); site opportunity from the codon model.",
 "   pN/pS << 1 = purifying selection.  [script 01]",
 "",
 "2. piTheta_wholepanel — pi/theta per site class, all 141. pi/theta = a_n*sum(pi_site)/S,",
 "   a_n(141)=5.522. The syn-vs-nonsyn contrast is a BETWEEN-clade frequency-spectrum signal.",
 "   [script 02]",
 "",
 "3. piTheta_withinclade — the same pi/theta within the Northern (n=77) and Southern (n=63)",
 "   clades. With clade structure removed, syn and NS pi/theta converge (~0.2): the whole-panel",
 "   gap is a clade-split effect, so within-species selection evidence rests on pN/pS. [script 03]",
 "",
 "'NA' = fewer than 2 sites, no CI. Complex II has no mtDNA-encoded subunits and is absent.",
 "Scripts: 01_pnps_bootstrap.py, 02_pi_theta_wholepanel.py, 03_pi_theta_withinclade.py, common.py.",
]
for i, t in enumerate(readme, 3):
    ws.cell(i, 1, t).font = (Font(name="Arial", bold=True, size=11)
                             if t[:2] in ("1.", "2.", "3.") else base)
ws.column_dimensions["A"].width = 100

sheet(wb.create_sheet("pN_pS_by_complex"),
      "pN/pS by OXPHOS complex (all 141)",
      "Primary purifying-selection statistic. Complex I CI does not overlap Complex III or IV -> Complex I under significantly relaxed constraint.",
      "pnps_bootstrap_by_complex.tsv")
sheet(wb.create_sheet("pN_pS_by_gene"),
      "pN/pS by gene (all 141)",
      "Every gene pN/pS << 1. COX1 most constrained; ND2/ND4 least.",
      "pnps_bootstrap_by_gene.tsv")
sheet(wb.create_sheet("piTheta_wholepanel_cx"),
      "pi/theta by complex, whole panel (syn vs nonsyn)",
      "Synonymous ~0.8 vs nonsynonymous ~0.32 (non-overlapping) = between-clade signal.",
      "pi_theta_wholepanel_by_complex.tsv")
sheet(wb.create_sheet("piTheta_wholepanel_gene"),
      "pi/theta by gene, whole panel (syn vs nonsyn)", None,
      "pi_theta_wholepanel_by_gene.tsv")
sheet(wb.create_sheet("piTheta_withinclade_cx"),
      "pi/theta by complex, WITHIN clade (North n=77, South n=63)",
      "Reviewer control: within each clade syn and NS converge (~0.2, overlapping CIs).",
      "pi_theta_withinclade_by_complex.tsv")
sheet(wb.create_sheet("piTheta_withinclade_gene"),
      "pi/theta by gene, WITHIN clade (North n=77, South n=63)", None,
      "pi_theta_withinclade_by_gene.tsv")

wb.save(opath("selection_stats_bootstrap_CIs.xlsx"))

# ============================ FIGURE ============================
fig, axes = plt.subplots(1, 3, figsize=(15, 5))

# Panel A: pN/pS by complex
rows = list(csv.DictReader(open(opath("pnps_bootstrap_by_complex.tsv")), delimiter="\t"))
labs = [r["Group"].replace("_", " ") for r in rows]
pe = [val(r["pN_pS"]) for r in rows]; lo = [val(r["CI_low"]) for r in rows]; hi = [val(r["CI_high"]) for r in rows]
x = np.arange(len(rows)); err = [np.array(pe)-np.array(lo), np.array(hi)-np.array(pe)]
cols = ["#c0392b" if l == "Overall" else "#2F5496" for l in labs]
axes[0].bar(x, pe, yerr=err, capsize=5, color=cols, edgecolor="black", linewidth=0.6)
axes[0].set_xticks(x); axes[0].set_xticklabels(labs, rotation=20, ha="right")
axes[0].axhline(1.0, ls="--", color="grey", lw=0.8)
axes[0].set_ylabel("pN/pS"); axes[0].set_title("A. pN/pS by complex (all 141)\n95% bootstrap CI", fontsize=11)
axes[0].set_ylim(0, max(hi)*1.25)

def grouped(ax, rows, keyfn, groups, combos, title):
    xb = np.arange(len(groups)); w = 0.8/len(combos)
    for k, (lab, color, pick) in enumerate(combos):
        pe = [val(pick(g)["pi_theta"]) for g in groups]
        lo = [val(pick(g)["CI_low"])   for g in groups]
        hi = [val(pick(g)["CI_high"])  for g in groups]
        err = [np.array(pe)-np.array(lo), np.array(hi)-np.array(pe)]
        off = (k - (len(combos)-1)/2) * w
        ax.bar(xb+off, pe, w, yerr=err, capsize=3, label=lab, color=color, edgecolor="black", linewidth=0.4)
    ax.set_xticks(xb); ax.set_xticklabels([g.replace("_", " ") for g in groups], rotation=20, ha="right")
    ax.set_ylabel("pi / theta"); ax.set_title(title, fontsize=11)

# Panel B: whole-panel syn vs NS
wp = list(csv.DictReader(open(opath("pi_theta_wholepanel_by_complex.tsv")), delimiter="\t"))
def wpick(cls): return lambda g: next(r for r in wp if r["Group"]==g and r["Class"]==cls)
groups = ["Complex_I","Complex_III","Complex_IV","Complex_V","Overall"]
grouped(axes[1], wp, None, groups,
        [("synonymous","#27ae60",wpick("synonymous")), ("nonsynonymous","#e67e22",wpick("nonsynonymous"))],
        "B. pi/theta whole panel\nsyn vs nonsyn (between-clade signal)")
axes[1].axhline(1.0, ls="--", color="grey", lw=0.8); axes[1].legend(fontsize=9)

# Panel C: within-clade syn vs NS, North & South
wc = list(csv.DictReader(open(opath("pi_theta_withinclade_by_complex.tsv")), delimiter="\t"))
def cpick(clade, cls): return lambda g: next(r for r in wc if r["Clade"]==clade and r["Group"]==g and r["Class"]==cls)
grouped(axes[2], wc, None, groups,
        [("North syn","#1f77b4",cpick("North","synonymous")), ("North non","#aec7e8",cpick("North","nonsynonymous")),
         ("South syn","#d62728",cpick("South","synonymous")), ("South non","#ff9896",cpick("South","nonsynonymous"))],
        "C. pi/theta WITHIN clade\nsyn~NS converge (control)")
axes[2].legend(fontsize=7, ncol=2)

plt.tight_layout()
plt.savefig(opath("selection_stats_figure.png"), dpi=200, bbox_inches="tight")
plt.savefig(opath("selection_stats_figure.pdf"), bbox_inches="tight")
print("04_make_workbook_figure.py: wrote selection_stats_bootstrap_CIs.xlsx and selection_stats_figure.png/.pdf")
