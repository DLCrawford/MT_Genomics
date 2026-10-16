#!/usr/bin/env python3
"""
05_per_individual_verification_table.py
---------------------------------------
Per-individual verification table of the synonymous / nonsynonymous variant
load, broken down by protein-coding gene. Two rows per individual (Syn, NonSyn);
141 individuals -> 282 data rows.

For each individual and gene, the value is the NUMBER OF SITES at which that
individual carries the derived (ALT) allele (genotype = 1), of the given class.
(This is an allele-occurrence count: a site carried by k fish contributes to k
individuals. It is NOT the number of segregating sites obs_S/obs_N, which counts
each site once regardless of carriers. A per-gene reconciliation of both is on
the 'gene_totals' sheet.)

Clade columns (both provided, because they differ):
  MITOTYPE       biological mt-clade call from sample metadata (N / S / A[admixed])
  Analysis_clade operational ALT-burden clade used by the pi-by-clade analysis
                 (north 77 / south 63 / ambiguous 1), from the membership file

Class of each site is taken from pi_by_clade_persite.tsv so labels match the
pi analysis exactly; sites absent there fall back to their SnpEff effect.

Inputs : ../vcf/haplotype_matrix.csv            (per-individual 0/1 genotypes)
         ../vcf/pi_by_clade_persite.tsv          (authoritative Syn/NonSyn class)
         ../vcf/pi_by_clade_persite.tsv.membership.tsv  (Analysis_clade, n_alt_panel)
         ../data_files_May/WGS_seq_plate.txt      (MITOTYPE)
Outputs: per_individual_SynNonSyn_by_gene.csv
         per_individual_SynNonSyn_by_gene.xlsx   (data + gene_totals + README sheets)
"""
import csv
from common import GENES, GENE2CX, vpath, opath
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MITO_SRC = vpath("../data_files_May/WGS_seq_plate.txt")

# ---- MITOTYPE ----
mito = {}
with open(MITO_SRC, encoding="latin-1") as f:
    for row in csv.DictReader(f, delimiter="\t"):
        mito[row["WGS_ID"].strip()] = row["MITOTYPE"].strip()

# ---- Analysis clade + ALT burden ----
aclade, nalt = {}, {}
with open(vpath("pi_by_clade_persite.tsv.membership.tsv")) as f:
    for row in csv.DictReader(f, delimiter="\t"):
        sid = row["Sample"].replace("_MT", "")
        aclade[sid] = row["Clade"]
        nalt[sid] = int(row["n_alt_panel"])

# ---- authoritative class per site (POS,REF,ALT) ----
site_class = {}
with open(vpath("pi_by_clade_persite.tsv"), encoding="latin-1") as f:
    r = csv.reader(f, delimiter="\t"); hdr = next(r); cols = hdr[:15]
    for row in r:
        if len(row) < 15: continue
        d = dict(zip(cols, row[:15]))
        cl = {"synonymous": "Syn", "nonsynonymous": "NonSyn"}.get(d["Class"])
        if cl: site_class[(d["POS"], d["REF"], d["ALT"])] = cl

def classify(pos, ref, alt, vt):
    c = site_class.get((pos, ref, alt))
    if c: return c
    return "NonSyn" if "missense" in vt else "Syn"   # fallback (stop_retained/splice -> Syn)

# ---- tally carriers per individual x gene x class ----
with open(vpath("haplotype_matrix.csv")) as f:
    rd = csv.DictReader(f)
    ids = [c[:-3] for c in rd.fieldnames if c.endswith("_MT")]
    # counts[ind][class][gene]
    counts = {i: {"Syn": {g: 0 for g in GENES}, "NonSyn": {g: 0 for g in GENES}} for i in ids}
    for r in rd:
        g = r["Gene"]; cls = classify(r["POS"], r["REF"], r["ALT"], r["Variant_type"])
        for i in ids:
            if r[i + "_MT"] == "1":
                counts[i][cls][g] += 1

def sortkey(i): 
    try: return (0, int(i))
    except: return (1, i)
ids_sorted = sorted(ids, key=sortkey)

meta = ["Individual", "MITOTYPE", "Analysis_clade", "n_alt_panel", "Class"]
header = meta + GENES + ["Total"]

# ---- CSV ----
with open(opath("per_individual_SynNonSyn_by_gene.csv"), "w", newline="") as o:
    w = csv.writer(o); w.writerow(header)
    for i in ids_sorted:
        for cls in ("Syn", "NonSyn"):
            row = [i, mito.get(i, "NA"), aclade.get(i, "NA"), nalt.get(i, "NA"), cls]
            vals = [counts[i][cls][g] for g in GENES]
            w.writerow(row + vals + [sum(vals)])

# ---- XLSX ----
wb = Workbook()
hf = Font(name="Arial", bold=True, color="FFFFFF", size=10)
fillN = PatternFill("solid", fgColor="2F5496")   # meta header
fillG = PatternFill("solid", fgColor="548235")   # gene header
base = Font(name="Arial", size=10)
synfill = PatternFill("solid", fgColor="EAF1FB")
nsfill  = PatternFill("solid", fgColor="FBEfe9")
thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center")

ws = wb.active; ws.title = "per_individual"
for j, h in enumerate(header, 1):
    c = ws.cell(1, j, h); c.font = hf; c.border = border; c.alignment = ctr
    c.fill = fillG if h in GENES else (PatternFill("solid", fgColor="7030A0") if h == "Total" else fillN)
r = 2
for i in ids_sorted:
    for cls in ("Syn", "NonSyn"):
        vals = [counts[i][cls][g] for g in GENES]
        rowvals = [int(i) if i.isdigit() else i, mito.get(i,"NA"), aclade.get(i,"NA"), nalt.get(i,"NA"), cls] + vals + [sum(vals)]
        for j, v in enumerate(rowvals, 1):
            c = ws.cell(r, j, v); c.font = base; c.border = border
            if j >= 4: c.alignment = ctr
            c.fill = synfill if cls == "Syn" else nsfill
        r += 1
ws.freeze_panes = "F2"
for j, h in enumerate(header, 1):
    ws.column_dimensions[get_column_letter(j)].width = 13 if j <= 5 else (6 if h in GENES else 7)

# gene_totals reconciliation sheet
ws2 = wb.create_sheet("gene_totals")
gt_hdr = ["Gene","Complex","Syn_allele_occurrences","NonSyn_allele_occurrences",
          "Syn_segregating_sites","NonSyn_segregating_sites"]
for j,h in enumerate(gt_hdr,1):
    c=ws2.cell(1,j,h); c.font=hf; c.fill=fillN; c.border=border; c.alignment=ctr
# segregating-site counts (each site once, if >=1 carrier)
seg = {g: {"Syn":0,"NonSyn":0} for g in GENES}
occ = {g: {"Syn":0,"NonSyn":0} for g in GENES}
for i in ids_sorted:
    for cls in ("Syn","NonSyn"):
        for g in GENES: occ[g][cls]+=counts[i][cls][g]
# recompute seg sites directly
with open(vpath("haplotype_matrix.csv")) as f:
    rd=csv.DictReader(f)
    for r_ in rd:
        g=r_["Gene"]; cls=classify(r_["POS"],r_["REF"],r_["ALT"],r_["Variant_type"])
        ac=sum(1 for i in ids if r_[i+"_MT"]=="1")
        if ac>=1: seg[g][cls]+=1
rr=2
for g in GENES:
    for j,v in enumerate([g,GENE2CX[g],occ[g]["Syn"],occ[g]["NonSyn"],seg[g]["Syn"],seg[g]["NonSyn"]],1):
        c=ws2.cell(rr,j,v); c.font=base; c.border=border
        if j>=3: c.alignment=ctr
    rr+=1
tot=[ "TOTAL","",sum(occ[g]["Syn"] for g in GENES),sum(occ[g]["NonSyn"] for g in GENES),
      sum(seg[g]["Syn"] for g in GENES),sum(seg[g]["NonSyn"] for g in GENES)]
for j,v in enumerate(tot,1):
    c=ws2.cell(rr,j,v); c.font=Font(name="Arial",bold=True,size=10); c.border=border
    if j>=3: c.alignment=ctr
for j in range(1,7): ws2.column_dimensions[get_column_letter(j)].width=max(12,len(gt_hdr[j-1])+2)

# README sheet
ws3 = wb.create_sheet("README")
notes=[
 "Per-individual synonymous / nonsynonymous variant load by gene",
 "",
 "Two rows per individual (Syn, NonSyn); 141 individuals = 282 data rows.",
 "Cell = number of SITES where that individual carries the derived (ALT) allele, of that class, in that gene.",
 "'Total' = row sum across the 13 genes.",
 "",
 "Clade columns differ by design:",
 "  MITOTYPE       - biological mt-clade from metadata: N=67, S=57, A(admixed)=17.",
 "  Analysis_clade - operational ALT-burden clade used by the pi-by-clade analysis:",
 "                   north 77 (<50 panel ALTs), south 63 (>200), ambiguous 1 (excluded from within-clade pi).",
 "  n_alt_panel    - that individual's total ALT count across the panel (the burden used for the split).",
 "",
 "Class labels are taken from pi_by_clade_persite.tsv so they match the pi analysis exactly.",
 "",
 "gene_totals sheet reconciles two different quantities:",
 "  allele_occurrences  = sum over individuals (this table's column totals); a k-carrier site counts k times.",
 "  segregating_sites   = each site counted once if >=1 carrier; compare to obs_S / obs_N in dnds_per_gene.tsv.",
]
for i,t in enumerate(notes,1):
    ws3.cell(i,1,t).font=(Font(name="Arial",bold=True,size=12) if i==1 else base)
ws3.column_dimensions["A"].width=110

wb.save(opath("per_individual_SynNonSyn_by_gene.xlsx"))
print("rows written:", len(ids_sorted)*2)
print("segregating-site totals  Syn=%d  NonSyn=%d" % (sum(seg[g]["Syn"] for g in GENES), sum(seg[g]["NonSyn"] for g in GENES)))
print("wrote per_individual_SynNonSyn_by_gene.csv and .xlsx")
